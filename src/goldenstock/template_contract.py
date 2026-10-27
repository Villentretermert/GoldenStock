"""Strict XLSX v1 contract. Parsing never evaluates workbook formulas."""
import ast
import io
import json
import math
import re
import zipfile
from lxml import etree
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill
from fastapi import HTTPException
from goldenstock.storage import validate_file

POLICY='gs017-v1'
HEADERS=['key','role','cell','unit','period','definition']
CELL=re.compile(r'\$?([A-Z]{1,2})\$?([1-9][0-9]{0,2})')


def invalid(message):
    raise HTTPException(422, message)


def formula_dependencies(formula):
    if not isinstance(formula,str) or not formula.startswith('=') or len(formula)>500:
        invalid('Invalid or oversized formula')
    body=formula[1:].replace("'Model'!",'').replace('Model!','')
    if '**' in body or '//' in body: invalid('Use Excel arithmetic syntax, including ^ for powers')
    if not re.fullmatch(r'[A-Z0-9$+\-*/^().\s]+',body):
        invalid('Functions, names and external references are not supported')
    refs=[]
    def reference(match):
        cell=match[1]+match[2]
        if cell not in refs: refs.append(cell)
        return 'v'+str(refs.index(cell))
    expression=CELL.sub(reference,body).replace('^','**')
    try: tree=ast.parse(expression.strip(),mode='eval')
    except (SyntaxError,RecursionError): invalid('Unsupported formula syntax')
    nodes=list(ast.walk(tree))
    if len(nodes)>100: invalid('Formula is too complex')
    allowed=(ast.Expression,ast.BinOp,ast.UnaryOp,ast.Add,ast.Sub,ast.Mult,ast.Div,
             ast.Pow,ast.UAdd,ast.USub,ast.Constant,ast.Name,ast.Load)
    for node in nodes:
        if not isinstance(node,allowed): invalid('Only arithmetic and mapped cell references are supported')
        if isinstance(node,ast.Name) and node.id not in {'v'+str(i) for i in range(len(refs))}:
            invalid('Functions, names and external references are not supported')
        if isinstance(node,ast.Constant) and (type(node.value) not in (int,float) or node.value not in (0,1)):
            invalid('Move numeric assumptions into mapped input cells; only 0 and 1 are formula constants')
    return refs


def validate_template(content):
    validate_file('template.xlsx',content,50*1024*1024)
    # Tighter expanded limits than generic documents; reject executable/link parts first.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos=archive.infolist()
        if len(infos)>250 or len({x.filename for x in infos})!=len(infos) or sum(x.file_size for x in infos)>5*1024*1024:
            invalid('Template package exceeds v1 limits')
        for item in infos:
            name=item.filename.lower()
            if any(x in name for x in ('vbaproject','externallinks','connections','querytables','macrosheets','embeddings','activex')):
                invalid('VBA, external connections and embedded objects are unsupported')
            if name.endswith(('.xml','.rels')):
                raw=archive.read(item)
                try:
                    tree=etree.fromstring(raw,etree.XMLParser(resolve_entities=False,no_network=True))
                    if tree.getroottree().docinfo.doctype: invalid('XML DTDs are unsupported')
                    if name.endswith('.rels') and any(n.get('TargetMode')=='External' for n in tree.iter()):
                        invalid('External links and hyperlinks are unsupported')
                except etree.XMLSyntaxError: invalid('Invalid workbook XML')
    try: book=load_workbook(io.BytesIO(content),data_only=False,read_only=False,keep_links=False)
    except Exception: invalid('Cannot parse XLSX template')
    try:
        if set(book.sheetnames)!={'Model','GS_Map'} or any(s.sheet_state!='visible' for s in book):
            invalid('Use exactly two visible sheets: Model and GS_Map')
        if book.defined_names or (book.calculation and book.calculation.iterate):
            invalid('Defined names and iterative calculations are unsupported')
        for sheet in book:
            if sheet.max_row>300 or sheet.max_column>30:
                invalid('Template exceeds 300 rows or 30 columns per sheet')
        mapping=book['GS_Map']
        if [mapping.cell(1,n).value for n in range(1,7)]!=HEADERS:
            invalid('GS_Map header must be: '+', '.join(HEADERS))
        rows=[];cells=set();keys=set();graph={}
        for values in mapping.iter_rows(min_row=2,max_col=6,values_only=True):
            if all(v is None for v in values): continue
            key,role,cell,unit,period,definition=values
            if not all(isinstance(v,str) and 0<len(v)<=160 for v in values):
                invalid('Every mapping field must be short text')
            if any(v.startswith('=') for v in values):
                invalid('Mapping metadata must be literal text, not formulas')
            if not re.fullmatch(r'[a-z][a-z0-9_]{0,49}',key) or key in keys:
                invalid('Mapping keys must be unique identifiers')
            if not re.fullmatch(r'[A-Z]{1,2}[1-9][0-9]{0,2}',cell) or cell in cells:
                invalid('Mapped cells must be unique A1 references')
            if role not in {'input','output','calculation'}:
                invalid('Mapping role must be input, output or calculation')
            if unit not in {'USD','EUR','GBP','RUB','ratio','multiple','count','years'}:
                invalid('Unsupported unit')
            if not re.fullmatch(r'(FY[0-9]{4}|assumption)',period): invalid('Use FYyyyy or assumption for period')
            target=book['Model'][cell]
            if target.row>300 or target.column>30: invalid('Mapped cell exceeds v1 limits')
            value=target.value
            if role=='input':
                if value is not None and (type(value) not in (int,float) or not math.isfinite(value)):
                    invalid('Input cells must be numeric or blank')
                graph[cell]=[]
            else: graph[cell]=formula_dependencies(value)
            rows.append(dict(zip(HEADERS,values),formula=value if role!='input' else None))
            cells.add(cell);keys.add(key)
        if len(rows)>100 or not any(r['role']=='input' for r in rows) or not any(r['role']=='output' for r in rows):
            invalid('Map 1–100 cells including inputs and outputs')
        for row in book['Model']:
            for cell in row:
                if cell.data_type=='f' and cell.coordinate not in cells: invalid('Every formula must be mapped')
        if any(ref not in cells for refs in graph.values() for ref in refs):
            invalid('Every referenced cell must be mapped')
        # Topological removal checks indirect cycles without evaluating any formula.
        pending=dict(graph);order=[]
        while pending:
            ready=[cell for cell,refs in pending.items() if all(r in order for r in refs)]
            if not ready: invalid('Circular references are unsupported')
            for cell in ready: order.append(cell);del pending[cell]
        return dict(policy=POLICY,fields=rows,order=order)
    finally: book.close()


def clean_workbook(contract):
    """Rebuild only the declared structure. Never copy private originals or cached values."""
    book=Workbook();sheet=book.active;sheet.title='Model'
    mapping=book.create_sheet('GS_Map');mapping.append(HEADERS)
    occupied={field['cell'] for field in contract['fields']}
    for field in contract['fields']:
        mapping.append([field[k] for k in HEADERS])
        cell=sheet[field['cell']]
        if cell.column>1 and sheet.cell(cell.row,cell.column-1).coordinate not in occupied:
            sheet.cell(cell.row,cell.column-1).value=field['key']
        cell.value=field['formula'] if field['role']!='input' else None
        cell.fill=PatternFill('solid',fgColor='E8F0FF' if field['role']=='input' else 'E1EDDF')
        cell.font=Font(name='Arial',color='0000FF' if field['role']=='input' else '193B30')
        cell.number_format='0.00% ' if field['unit']=='ratio' else '#,##0.00;(#,##0.00);"-"'
    for col in 'ABCDEF': mapping.column_dimensions[col].width=26
    sheet.column_dimensions['A'].width=28
    for col in 'BCDEFG': sheet.column_dimensions[col].width=20
    stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
