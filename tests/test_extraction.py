"""Synthetic corpus tests; these do not estimate quality on real issuer reports."""
import io
import os
from pathlib import Path
import shutil
import pytest
from goldenstock.extraction import extract, ExtractionError
from tests.test_workspace import workspace, upload
from scripts.extraction_fixtures import simple_pdf, workbook_bytes






def test_xlsx_preserves_negative_blank_merge_and_formula(tmp_path):
    path=tmp_path/'book.xlsx'; path.write_bytes(workbook_bytes())
    result=extract(path,'.xlsx')
    cells={i['location']['cell']:i for i in result['items']}
    assert cells['B2']['text']=='-120'
    assert cells['C2']['text'] is None
    assert cells['B3']['location']['merged_range']=='A3:C3'
    assert cells['B3']['text'] is None
    assert cells['B4']['formula']=='=B2*2' and cells['B4']['text'] is None
    assert 'formula_cache_unverified' in result['warnings']
    assert all(i['review_status']=='unreviewed' and i['confidence'] is None for i in result['items'])


def test_pdf_text_page_and_footnote(tmp_path):
    path=tmp_path/'text.pdf'; path.write_bytes(simple_pdf())
    result=extract(path,'.pdf')
    assert 'Profit: -120' in result['items'][0]['text']
    assert 'Footnote 1' in result['items'][0]['text']
    assert result['items'][0]['location']['page']==1


def test_pdf_table_cells(tmp_path):
    stream=b'50 600 m 350 600 l 350 700 l 50 700 l h S 50 650 m 350 650 l S 200 600 m 200 700 l S '
    stream+=b'BT /F1 16 Tf 60 670 Td (Profit) Tj 150 0 Td (-120) Tj ET'
    path=tmp_path/'table.pdf'; path.write_bytes(simple_pdf(stream))
    cells=[i for i in extract(path,'.pdf')['items'] if i['kind']=='cell']
    assert len(cells)==4
    assert cells[1]['text']=='-120' and cells[1]['location']['column']==2
    assert cells[1]['location']['cell_bbox'] is not None


def test_pptx_table_cells(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
    table=slide.shapes.add_table(2,2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    table.cell(0,0).text='Profit'; table.cell(0,1).text='(120)'
    table.cell(1,0).merge(table.cell(1,1)); table.cell(1,0).text='Footnote 1'
    path=tmp_path/'deck.pptx'; deck.save(path)
    items=extract(path,'.pptx')['items']
    assert items[1]['text']=='(120)' and items[1]['location']['slide']==1
    assert items[2]['merge_origin'] and items[3]['merged_follower']


def test_scan_ocr(tmp_path):
    if not shutil.which('tesseract'):
        pytest.skip('Real Tesseract test runs in Docker')
    from PIL import Image, ImageDraw, ImageFont
    image=Image.new('RGB',(1200,400),'white')
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40)
    ImageDraw.Draw(image).text((60,80),'Profit -120 USD million',font=font,fill='black')
    path=tmp_path/'scan.pdf'; image.save(path,'PDF',resolution=150)
    result=extract(path,'.pdf')
    assert 'Profit' in ' '.join(i['text'] for i in result['items'])
    assert all(i['method']=='ocr' and 0<=i['confidence']<=1 for i in result['items'])


def test_extraction_queue_privacy_and_idempotency(workspace, monkeypatch):
    from goldenstock.jobs import process_one
    from goldenstock.models import Extraction
    from sqlalchemy.orm import Session
    from sqlalchemy import select, func
    from goldenstock.ui import mount_ui
    app,connect=workspace; mount_ui(app)
    # subprocess needs an explicit module path in native pytest, unlike Docker.
    monkeypatch.setenv('PYTHONPATH',str(Path('src').resolve()))
    a,b,e=connect(),connect('bob'),connect('eve')
    receipt=upload(a,content=workbook_bytes(),name='test.xlsx').json(); job=receipt['job_id']
    url=f'/api/v1/jobs/{job}/extract'
    assert a.post(url).status_code==409
    process_one(app.state.engine,app.state.settings)
    assert a.post(url).status_code==200
    assert a.post(url).status_code==200
    process_one(app.state.engine,app.state.settings)
    result=a.get(f'/api/v1/jobs/{job}/extraction')
    assert result.status_code==200, a.get(f'/api/v1/jobs/{job}').text
    assert result.json()['source_sha256']==receipt['sha256']
    assert a.get(f'/extraction/{job}').status_code==200
    assert a.get(f'/api/v1/jobs/{job}/extraction?limit=201').status_code==422
    for client in (b,e):
        assert client.get(f'/api/v1/jobs/{job}/extraction').status_code==404
        assert client.post(url).status_code==404
        assert client.get(f'/extraction/{job}').status_code==404
    assert a.post(url).status_code==200
    process_one(app.state.engine,app.state.settings)
    with Session(app.state.engine) as db:
        assert db.scalar(select(func.count()).select_from(Extraction))==1


def test_corrupt_document_fails_without_result(workspace, monkeypatch):
    from goldenstock.jobs import process_one
    app,connect=workspace; client=connect()
    monkeypatch.setenv('PYTHONPATH',str(Path('src').resolve()))
    job=upload(client).json()['job_id']
    process_one(app.state.engine,app.state.settings)
    client.post(f'/api/v1/jobs/{job}/extract')
    process_one(app.state.engine,app.state.settings)
    status=client.get(f'/api/v1/jobs/{job}').json()
    assert status['state']=='failed' and status['error_code']=='document_parse_failed'
    assert client.get(f'/api/v1/jobs/{job}/extraction').status_code==409


def test_dimension_limit(tmp_path):
    from openpyxl import Workbook
    book=Workbook(); book.active['A50001']='Sparse oversized sheet'
    path=tmp_path/'large.xlsx'; book.save(path)
    with pytest.raises(ExtractionError,match='extraction_limit'):
        extract(path,'.xlsx')


@pytest.mark.parametrize('change',['revoke','replace_lease'])
def test_results_not_committed_after_authority_changes(workspace,monkeypatch,change):
    from goldenstock import jobs
    from goldenstock.models import Job, User, Extraction
    from sqlalchemy.orm import Session
    app,connect=workspace; client=connect()
    receipt=upload(client,content=workbook_bytes(),name='test.xlsx').json()
    jobs.process_one(app.state.engine,app.state.settings)
    client.post(f'/api/v1/jobs/{receipt["job_id"]}/extract')
    leased=jobs.claim(app.state.engine)
    assert jobs.renew(app.state.engine,app.state.settings,*leased)
    def parser(*args):
        with Session(app.state.engine) as db,db.begin():
            job=db.get(Job,leased[0])
            if change=='revoke': db.get(User,job.actor_id).active=False
            else: job.lease_token='replacement'
        return {'parser_version':'synthetic','items':[]}
    monkeypatch.setattr(jobs,'run_parser',parser)
    jobs.execute(app.state.engine,app.state.settings,leased)
    with Session(app.state.engine) as db:
        assert db.get(Extraction,receipt['version_id']) is None
    assert not jobs.renew(app.state.engine,app.state.settings,*leased)
