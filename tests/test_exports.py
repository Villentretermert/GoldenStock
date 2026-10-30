"""Known-value parity, immutable evidence, private downloads and safe editable workbooks."""
import copy
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import json
import openpyxl
import pdfplumber
from pptx import Presentation
import pytest
from sqlalchemy.orm import Session
from test_workspace import workspace
from test_valuations import source, api_payload, create, numbers
from test_reviews import decide
from goldenstock.models import ValuationRun
from goldenstock.export_documents import wrapped_lines, text_pages
from goldenstock.export_workbook import expression
from goldenstock.calculation_engine import parse_formula


def downloaded(client, rid, kind):
    response = client.get(f'/api/v1/valuations/{rid}/export/{kind}')
    assert response.status_code == 200, response.text[:500] if response.status_code != 200 else ''
    assert response.headers['cache-control'] == 'no-store'
    assert hashlib.sha256(response.content).hexdigest() == response.headers['x-content-sha256']
    assert 'attachment;' in response.headers['content-disposition']
    return response.content


def save_fixture(name, content):
    folder = Path('.build-export-fixtures'); folder.mkdir(exist_ok=True)
    (folder / name).write_bytes(content)


def pdf_text(raw):
    with pdfplumber.open(BytesIO(raw)) as pdf: return '\n'.join(p.extract_text() or '' for p in pdf.pages)


def pptx_text(raw):
    deck = Presentation(BytesIO(raw)); texts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame: texts.append(shape.text)
            if shape.has_table:
                texts += [c.text for row in shape.table.rows for c in row.cells]
    return '\n'.join(texts)


@pytest.mark.parametrize('method,expected', [('pe',90), ('pb',120), ('ddm',100), ('fcfe',100)])
def test_all_formats_match_saved_values_and_provenance(workspace, method, expected):
    app, connect = workspace; client = connect(); cid, _ = source(client, app)
    data = api_payload(cid, method)
    data.update(operating=True, operating_month='2025-12', operating_note='Synthetic monthly cohort with constant revenue and retention',
                scenarios={'base':numbers(method, True)})
    rid = create(client, data).json()['id']; saved = client.get('/api/v1/valuations/'+rid).json()
    raw = {kind:downloaded(client,rid,kind) for kind in ('xlsx','pdf','pptx')}
    book = openpyxl.load_workbook(BytesIO(raw['xlsx']), data_only=True)
    assert book['Summary']['B8'].value == pytest.approx(expected)
    formulas = openpyxl.load_workbook(BytesIO(raw['xlsx']))
    assert formulas.calculation.fullCalcOnLoad and formulas.calculation.calcMode == 'auto'
    assert formulas['Summary']['B8'].data_type == 'f'
    assert len(book['Sensitivity']['A']) == (16 if method in ('ddm','fcfe') else 10)
    evidence = '\n'.join(str(c.value or '') for row in book['Evidence'] for c in row)
    for content in (evidence, pdf_text(raw['pdf']), pptx_text(raw['pptx'])):
        assert f'{expected:.2f}' in content
        assert saved['result_hash'] in content.replace('\n','')
        assert 'gs020-v1' in content
        assert 'Synthetic explicit assumption' in content
        assert 'Manual inputs' in content
    # Synthetic artifacts are local QA outputs only, never committed or sent to a service.
    folder = Path('.build-export-fixtures'); folder.mkdir(exist_ok=True)
    for kind, value in raw.items(): (folder / f'{method}.{kind}').write_bytes(value)
    (folder / f'{method}.json').write_text(json.dumps(saved),encoding='utf-8')


def test_exports_acl_stale_review_and_tampered_receipt(workspace):
    app, connect = workspace; a, b, e = connect(), connect('bob'), connect('eve')
    cid, fid = source(a, app, True); rid = create(a, api_payload(cid)).json()['id']
    url = f'/api/v1/valuations/{rid}/export/'
    for other in (b,e):
        for kind in ('xlsx','pdf','pptx'): assert other.get(url+kind).status_code == 404
    assert a.get(url+'zip').status_code == 422
    assert 'Source review current' in pdf_text(downloaded(a, rid, 'pdf'))
    assert decide(a, fid, 'unreviewed', 1).status_code == 201
    for kind in ('pdf','pptx'):
        content = (pdf_text if kind=='pdf' else pptx_text)(downloaded(a,rid,kind))
        assert 'HISTORICAL' in content and 'source review changed' in content
        assert 'accepted' in content  # Historical provenance remains as captured.
    with Session(app.state.engine) as db, db.begin(): db.get(ValuationRun,rid).output='{}'
    for kind in ('xlsx','pdf','pptx'): assert a.get(url+kind).status_code==409


def test_export_source_formula_binding_and_text_injection(workspace):
    app, connect = workspace; a=connect();cid,_=source(a,app)
    data=api_payload(cid);data['title']='=HYPERLINK("https://example.invalid", "Title")'
    data['method_note']='<img src="https://example.invalid/pixel">' + 'Q'*1900
    data['scenarios']['base']['profit_owners']=dict(source_key='profit_owners',note='Source profit checked for parent scope',confirmed=True)
    rid=create(a,data).json()['id']; raw=downloaded(a,rid,'xlsx')
    book=openpyxl.load_workbook(BytesIO(raw))
    assert book['Summary']['A1'].data_type=='s' and book['Summary']['A1'].value==data['title']
    assert 'Calc_base' in book['Val_base']['C9'].value
    assert not book._external_links
    with ZipFile(BytesIO(raw)) as z:
        assert not any('vba' in n.lower() or 'externalLink' in n for n in z.namelist())
    pdf=downloaded(a,rid,'pdf');slides=downloaded(a,rid,'pptx')
    assert 'Q'*100 in pdf_text(pdf).replace('\n','')
    assert 'Q'*100 in pptx_text(slides).replace('\n','')
    save_fixture('long-text.pptx', slides)


def test_blocked_valuation_cannot_export_and_ui_downloads(workspace):
    app,connect=workspace;a=connect();cid,_=source(a,app);data=api_payload(cid)
    data['scenarios']['base']['shares']['value']='0'
    rid=create(a,data).json()['id']
    for kind in ('xlsx','pdf','pptx'): assert a.get(f'/api/v1/valuations/{rid}/export/{kind}').status_code==409
    from goldenstock.ui import mount_ui
    mount_ui(app)
    assert 'Download XLSX' not in a.get('/valuations/'+rid).text
    rid=create(a,api_payload(cid),'valid-export-ui').json()['id']
    page=a.get('/valuations/'+rid).text
    for kind in ('XLSX','PDF','PPTX'): assert 'Download '+kind in page


def test_excel_compiler_preserves_unary_power_and_long_source_text():
    formula='=-B1^B2^B3'
    assert expression(parse_formula(formula)[0], {'B1':'C8','B2':'C9','B3':'C10'})=='(((-C8)^C9)^C10)'
    original='X'*5000
    assert ''.join(wrapped_lines([original]))==original
    receipt = 'Valuation result SHA-256: ' + 'abcdef0123456789'*4
    pages = list(text_pages(['One', 'Two', receipt], limit=3))
    assert any(receipt in ''.join(page) for page in pages)


def test_three_scenarios_source_links_quote_and_invalid_sensitivity(workspace):
    from test_templates import post
    from test_calculations import payload, run
    app,connect=workspace;a=connect();calc=payload(post(a).json())
    calc['scenarios']['optimistic']=copy.deepcopy(calc['scenarios']['base'])
    calc['scenarios']['optimistic']['profit_owners']['value']='200'
    calc['scenarios']['stress']=copy.deepcopy(calc['scenarios']['base'])
    calc['scenarios']['stress']['profit_owners']['value']='50'
    cid=run(a,calc).json()['id'];data=api_payload(cid)
    data['quote']=dict(price='60',currency='USD',date='2025-12-31',instrument='Synthetic common share',kind='ordinary_share',source='Synthetic quote for testing')
    data['scenarios']={name:numbers() for name in calc['scenarios']}
    for inputs in data['scenarios'].values():
        inputs['profit_owners']=dict(source_key='profit_owners',note='Parent profit source checked',confirmed=True)
    rid=create(a,data).json()['id'];raw=downloaded(a,rid,'xlsx')
    book=openpyxl.load_workbook(BytesIO(raw),data_only=True)
    assert [book['Summary'][f'B{i}'].value for i in range(8,11)]==[90,190,40]
    assert book['Summary']['D8'].value==.5
    for kind,extract in [('pdf',pdf_text),('pptx',pptx_text)]:
        content=extract(downloaded(a,rid,kind))
        assert all(str(v)+'.00' in content for v in (90,190,40))
    save_fixture('linked.xlsx', raw)
    dcf=api_payload(cid,'ddm');dcf['scenarios']['base']['cost_equity']['value']='0.005'
    rid=create(a,dcf,'invalid-sensitivity-export').json()['id']
    raw=downloaded(a,rid,'xlsx');book=openpyxl.load_workbook(BytesIO(raw),data_only=True)
    assert book['Sensitivity']['D8'].value=='Blocked'
    assert 'Blocked' in pdf_text(downloaded(a,rid,'pdf'))
    save_fixture('boundary.xlsx', raw)
