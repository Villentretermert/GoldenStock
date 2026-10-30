"""Local Python report writers; no HTML loading, external assets or remote services."""
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from goldenstock.export_content import clean, sections, summary_table, sensitivity_tables, review_label

INK = '193B30'


@lru_cache(maxsize=1)
def fonts():
    pairs = [('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
             ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf')]
    for regular, bold in pairs:
        if Path(regular).is_file() and Path(bold).is_file():
            pdfmetrics.registerFont(TTFont('GS', regular)); pdfmetrics.registerFont(TTFont('GSBold', bold))
            return 'Arial' if regular.startswith('C:') else 'DejaVu Sans'
    raise RuntimeError('Install DejaVu Sans fonts for local report rendering')


def pdf_bytes(report):
    fonts(); stream = BytesIO()
    body = ParagraphStyle('body', fontName='GS', fontSize=9, leading=14, spaceAfter=5, splitLongWords=True)
    title = ParagraphStyle('title', parent=body, fontName='GSBold', fontSize=22, leading=28, textColor=colors.HexColor('#' + INK), spaceAfter=14)
    heading = ParagraphStyle('heading', parent=body, fontName='GSBold', fontSize=13, leading=18, spaceBefore=12, spaceAfter=8, keepWithNext=True)
    def p(value, style=body): return Paragraph(escape(clean(value)).replace('\n', '<br/>'), style)
    story = [p(report['title'], title), p(report['company'] + ' | ' + report['input']['method'].upper()),
             p('Valuation date: ' + report['input']['valuation_date'] + ' | ' + report['input']['currency'] + ' whole units'),
             p(review_label(report)), Spacer(1, 12)]

    def table(rows):
        # Table width is fixed; text wraps inside each cell.
        t = Table([[p(c) for c in row] for row in rows], colWidths=[(A4[0] - 96) / len(rows[0])] * len(rows[0]), repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EDF3EF')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 7), ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor('#' + INK))]))
        return t
    story.append(table(summary_table(report)))
    for name, rows in sensitivity_tables(report): story += [p(name, heading), table(rows)]
    for name, lines in sections(report):
        group = [p(name, heading)] + [p(line) if line else Spacer(1, 8) for line in lines]
        # Keep short method/receipt blocks intact; source appendices paginate naturally.
        if name == 'Version and provenance' or sum(len(line) for line in lines) < 900:
            story.append(KeepTogether(group))
        else: story.extend(group)
    footer_style = ParagraphStyle('footer', parent=body, fontSize=7, leading=10)
    def footer(canvas, doc):
        footer = p(f'{report["export_version"]} | {report["exported_at"]} | Page {doc.page}\nValuation {report["id"]}', footer_style)
        _, height = footer.wrap(A4[0] - 96, 35); footer.drawOn(canvas, 48, 22)
    SimpleDocTemplate(stream, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=40, bottomMargin=62,
                      title=clean(report['title']), author='GoldenStock').build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def wrapped_lines(lines, width=850, size=18):
    """Measure every line; split long unbroken source tokens without dropping characters."""
    fonts(); result = []
    for line in lines:
        for paragraph in clean(line).split('\n'):
            while paragraph:
                end = len(paragraph)
                if pdfmetrics.stringWidth(paragraph, 'GS', size) > width:
                    lo, hi = 1, end
                    while lo < hi:
                        mid = (lo + hi + 1) // 2
                        if pdfmetrics.stringWidth(paragraph[:mid], 'GS', size) <= width: lo = mid
                        else: hi = mid - 1
                    end = lo
                    space = paragraph.rfind(' ', 0, end)
                    if space > end // 2: end = space + 1
                result.append(paragraph[:end]); paragraph = paragraph[end:]
            result.append('')
    while result and not result[-1]: result.pop()
    return result


def pptx_bytes(report):
    deck = Presentation(); deck.slide_width = 12192000; deck.slide_height = 6858000
    family = fonts()  # Use the same font metrics for wrapping and slide text.
    deck.core_properties.title = clean(report['title']); deck.core_properties.author = 'GoldenStock'
    def box(slide, value, x, y, w, h, size=18, bold=False, color=INK):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for index, line in enumerate(clean(value).split('\n')):
            para = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
            para.text = line; para.font.name = family; para.font.size = Pt(size); para.font.bold = bold
            para.font.color.rgb = RGBColor.from_string(color); para.line_spacing = Pt(size * 1.12)
            para.space_before = Pt(0); para.space_after = Pt(0)
        return shape
    def slide(title):
        s = deck.slides.add_slide(deck.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string('F7F8F4')
        box(s, title, .55, .4, 12.2, .9, 28, True)
        box(s, f'{report["export_version"]} | {report["exported_at"]} | {report["id"]} | {len(deck.slides)}', .55, 7.08, 12.2, .2, 9, color='586B61')
        return s
    s = slide('Common equity valuation')
    title_lines = wrapped_lines([report['title']], width=850, size=24)
    box(s, '\n'.join(title_lines).rstrip(), .55, 1.45, 12.2, 1.4, 24, True)
    box(s, f'{report["company"]} | {report["method"].upper()}\nAs of {report["input"]["valuation_date"]} | {report["input"]["currency"]} whole units', .55, 3.05, 12.2, 1.2, 22)
    box(s, '\n'.join(wrapped_lines([review_label(report)], size=18)).rstrip(), .55, 5.05, 12.2, 1.0, 18, color='8A5518')
    def table(title, rows):
        s = slide(title); count = len(rows)
        table = s.shapes.add_table(count, len(rows[0]), Inches(.55), Inches(1.6), Inches(12.2), Inches(min(4.85, count * .52))).table
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                cell = table.cell(i, j); cell.text = clean(value); cell.margin_left = cell.margin_right = Inches(.12)
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(INK if i == 0 else 'EDF3EF')
                for para in cell.text_frame.paragraphs:
                    para.font.name = family; para.font.size = Pt(16); para.font.bold = i == 0
                    para.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else INK)
    table('Scenario comparison (' + report['input']['currency'] + ')', summary_table(report))
    for name, rows in sensitivity_tables(report): table(name, rows)
    for title, lines in sections(report):
        lines = wrapped_lines(lines)
        for start in range(0, len(lines), 16):
            s = slide(title + (' (continued)' if start else ''))
            box(s, '\n'.join(lines[start:start + 16]).rstrip(), .55, 1.55, 12.2, 5.15)
    stream = BytesIO(); deck.save(stream); return stream.getvalue()
