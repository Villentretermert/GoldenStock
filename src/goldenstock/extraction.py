"""Local raw extraction. Values are evidence, never approved financial facts."""
import json
from pathlib import Path

PARSER_VERSION = 'gs010-v1'
MAX_ITEMS = 50000
MAX_BYTES = 8 * 1024 * 1024
MAX_PAGES = 250


class ExtractionError(Exception):
    pass


class Result:
    def __init__(self):
        self.items = []
        self.warnings = set()
        self.size = 0

    def add(self, kind, text, location, method='native', confidence=None, **extra):
        item = dict(kind=kind, text=None if text is None else str(text), location=location,
                    method=method, confidence=confidence, review_status='unreviewed', **extra)
        self.size += len(json.dumps(item, ensure_ascii=True).encode())
        if len(self.items) >= MAX_ITEMS or self.size > MAX_BYTES - 10000:
            raise ExtractionError('extraction_limit')
        self.items.append(item)

    def output(self):
        return dict(parser_version=PARSER_VERSION, items=self.items, warnings=sorted(self.warnings),
                    language='en', language_status='analyst_confirmation_required',
                    review_status='unreviewed')


def extract_xlsx(path, result):
    from openpyxl import load_workbook
    # Never evaluate formulas or follow external links; retain cached values separately.
    with open(path, 'rb') as stream, open(path, 'rb') as cached_stream:
        book = load_workbook(stream, data_only=False, keep_links=False)
        cached = load_workbook(cached_stream, data_only=True, keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row * sheet.max_column > MAX_ITEMS:
                    raise ExtractionError('extraction_limit')
                for row in sheet.iter_rows():
                    for cell in row:
                        formula = cell.data_type == 'f'
                        # Range membership preserves blank merged followers, not duplicated values.
                        merged = next((str(r) for r in sheet.merged_cells.ranges if cell.coordinate in r), None)
                        value = cached[sheet.title][cell.coordinate].value if formula else cell.value
                        if formula:
                            result.warnings.add('formula_cache_unverified')
                        result.add('cell', value, dict(sheet=sheet.title, cell=cell.coordinate,
                            row=cell.row, column=cell.column, merged_range=merged),
                            formula=str(cell.value) if formula else None, number_format=cell.number_format,
                            hidden_sheet=sheet.sheet_state != 'visible')
        finally:
            book.close(); cached.close()


def extract_pptx(path, result):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    with open(path, 'rb') as stream:
        deck = Presentation(stream)
        if len(deck.slides) > MAX_PAGES:
            raise ExtractionError('extraction_limit')
        def shapes(items, slide, prefix=''):
            for shape in items:
                key = prefix + str(shape.shape_id)
                loc = dict(slide=slide, shape=key)
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, slide, key + '/')
                elif shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            result.add('cell', cell.text, dict(loc, table=key, row=r, column=c),
                                merge_origin=cell.is_merge_origin, merged_follower=cell.is_spanned,
                                span_width=cell.span_width, span_height=cell.span_height)
                elif shape.has_text_frame:
                    result.add('text', shape.text, loc)
                else:
                    result.warnings.add('nontext_slide_content_not_extracted')
        for i, slide in enumerate(deck.slides, 1):
            shapes(slide.shapes, i)


def extract_pdf(path, result):
    import pdfplumber
    import pypdfium2 as pdfium
    import pytesseract
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise ExtractionError('extraction_limit')
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ''
            if text.strip():
                result.add('text', text, dict(page=i))
                for t, table in enumerate(page.find_tables(), 1):
                    for r, row in enumerate(table.extract(), 1):
                        for c, value in enumerate(row, 1):
                            result.add('cell', value, dict(page=i, table=t, row=r, column=c,
                                table_bbox=list(table.bbox), cell_bbox=table.rows[r-1].cells[c-1]))
            # OCR image-bearing pages too: a header alone does not prove the body has text.
            if page.images or not text.strip():
                if page.width * page.height * 9 > 25000000:
                    raise ExtractionError('extraction_limit')
                with pdfium.PdfDocument(path) as raster:
                    image_page = raster[i-1]
                    bitmap = image_page.render(scale=3)
                    image = bitmap.to_pil()
                    try:
                        data = pytesseract.image_to_data(image, lang='eng', timeout=30,
                            output_type=pytesseract.Output.DICT)
                    finally:
                        image.close(); bitmap.close(); image_page.close()
                result.warnings.add('ocr_layout_requires_review')
                if text.strip():
                    result.warnings.add('native_and_ocr_overlap')
                for n, word in enumerate(data['text']):
                    if word.strip():
                        result.add('word', word, dict(page=i, bbox_pixels=[data[k][n] for k in ('left','top','width','height')],
                            render_scale=3, block=data['block_num'][n], paragraph=data['par_num'][n],
                            line=data['line_num'][n], word=data['word_num'][n]),
                            method='ocr', confidence=max(0, float(data['conf'][n])) / 100)
            page.close()


def extract(path, extension):
    result = Result()
    parsers = {'.pdf': extract_pdf, '.xlsx': extract_xlsx, '.pptx': extract_pptx}
    if extension not in parsers:
        raise ExtractionError('unsupported_format')
    parsers[extension](path, result)
    if not result.items:
        raise ExtractionError('no_extractable_content')
    return result.output()


def main():
    import sys
    # Linux container is the supported production parser runtime.
    if sys.platform == 'linux':
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (3 * 1024**3, 3 * 1024**3))
        resource.setrlimit(resource.RLIMIT_CPU, (180, 180))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_BYTES, MAX_BYTES))
    try:
        payload = extract(sys.argv[1], sys.argv[2])
    except ExtractionError as exc:
        payload = {'error': str(exc)}
    except Exception:
        # Document contents and library errors must not leak into worker logs.
        payload = {'error': 'document_parse_failed'}
    sys.stdout.write(json.dumps(payload, ensure_ascii=True))


if __name__ == '__main__':
    main()
