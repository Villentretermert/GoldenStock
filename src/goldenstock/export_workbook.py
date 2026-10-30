"""Portable Excel formulas with server-result caches and explicit precision limits."""
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties
from goldenstock.calculation_engine import parse_formula
from goldenstock.export_content import clean, review_label, sections

GREEN = '193B30'
PALE = 'EDF3EF'
NUM = '#,##0.00;(#,##0.00);"-"'


def text(cell, value):
    # User strings, including =/+/@ prefixes, can never become workbook formulas.
    cell.value = clean(value)
    cell.data_type = 's'


def expression(node, refs):
    if node[0] == 'ref': return refs[node[1]]
    if node[0] == 'number': return node[1]
    if node[0] == 'unary': return '(' + node[1] + expression(node[2], refs) + ')'
    return '(' + expression(node[2], refs) + node[1] + expression(node[3], refs) + ')'


def domain(method, refs, operating):
    """Excel guards mirror the economic domain; spreadsheet errors stay errors, never zero."""
    r = refs
    terms = [f'{r["shares"]}>0']
    if method in ('ddm', 'fcfe'):
        terms += [f'{r["cost_equity"]}>0', f'{r["cost_equity"]}<=1', f'{r["terminal_growth"]}>-1',
                  f'{r["terminal_growth"]}<{r["cost_equity"]}', f'{r["cash_5"]}>0']
        if method == 'ddm': terms += [r['cash_' + str(n)] + '>=0' for n in range(1, 6)]
    elif method == 'pe':
        terms += [f'{r["preferred_dividends"]}>=0', f'{r["profit_owners"]}-{r["preferred_dividends"]}>0', f'{r["multiple"]}>0']
    else:
        terms += [f'{r["equity_total"]}-{r["equity_nci"]}-{r["equity_preferred"]}>0', f'{r["equity_preferred"]}>=0', f'{r["multiple"]}>0']
    if operating:
        terms += [f'{r["mau"]}>0', f'{r["cohort_start"]}>0', f'{r["cohort_retained"]}>=0',
                  f'{r["cohort_retained"]}<{r["cohort_start"]}', f'{r["monthly_revenue"]}>=0']
        for k in ('mau', 'cohort_start', 'cohort_retained'): terms.append(f'MOD({r[k]},1)=0')
        for k in ('contribution_margin', 'monthly_discount'): terms += [f'{r[k]}>=0', f'{r[k]}<=1']
    return terms


def workbook_bytes(report):
    wb = Workbook(); wb.remove(wb.active); caches = {}
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    data = report['input']; source = data['source']['input']; result = report['result']

    def sheet(name, title):
        ws = wb.create_sheet(name); ws.sheet_view.showGridLines = False
        ws.merge_cells('A1:E1'); text(ws['A1'], title); ws['A1'].font = Font(name='Arial', size=18, bold=True, color=GREEN)
        ws.row_dimensions[1].height = 36
        text(ws['A2'], report['export_version'] + ' | ' + report['exported_at']); ws.merge_cells('A2:E2')
        ws.freeze_panes = 'C8'
        for col, width in {'A': 30, 'B': 24, 'C': 26, 'D': 90, 'E': 20}.items(): ws.column_dimensions[col].width = width
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.print_title_rows = '1:7'
        ws.oddFooter.center.text = 'GoldenStock | &P / &N'
        return ws

    def header(ws, row, values):
        for col, value in enumerate(values, 1):
            c = ws.cell(row, col); text(c, value)
            c.fill = PatternFill('solid', fgColor=GREEN); c.font = Font(name='Arial', color='FFFFFF', bold=True)
            c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[row].height = 32

    def formula(ws, cell, value, cached):
        ws[cell] = value; caches[(ws.title, cell)] = cached
        ws[cell].font = Font(name='Arial', color='008000' if '!' in value else '000000', size=11)

    def number_format(value, unit=''):
        return '0.00E+00' if abs(Decimal(value)) >= Decimal('1e15') else ('0.00%' if unit == 'ratio' else NUM)

    def numeric(ws, cell, value, unit='', editable=True):
        ws[cell] = float(Decimal(value))
        ws[cell].font = Font(name='Arial', color='0000FF' if editable else GREEN, size=11)
        ws[cell].number_format = number_format(value, unit)

    summary = sheet('Summary', report['title'])
    summary.column_dimensions['D'].width = 32
    for row, label in [(3, report['company']), (4, data['method'].upper() + ' | ' + data['currency'] + ' whole units'),
                       (5, 'Valuation date: ' + data['valuation_date'])]: text(summary.cell(row, 1), label)
    header(summary, 7, ['Scenario', 'Per common share', 'Common equity', 'Upside vs quote', 'Input domain'])
    build_refs = {}; val_refs = {}
    for name, inputs in source['scenarios'].items():
        ws = sheet('Calc_' + name, 'Source calculation: ' + name)
        fields = source['template']['contract']['fields']
        refs = {f['cell']: 'C' + str(i + 8) for i, f in enumerate(fields)}
        keys = {f['key']: refs[f['cell']] for f in fields}; build_refs[name] = keys
        text(ws['A3'], 'All inputs numeric and bounded'); text(ws['A4'], 'Control gate (TRUE if none)')
        input_refs = [keys[f['key']] for f in fields if f['role'] == 'input']
        formula(ws, 'B3', '=AND(' + ','.join(f'AND(ISNUMBER({r}),ABS({r})<=1E30)' for r in input_refs) + ')', True)
        header(ws, 7, ['Field', 'Unit', 'Value / formula', 'Definition', 'Period'])
        saved = data['source']['result']['scenarios'][name]
        for i, f in enumerate(fields, 8):
            for col, value in [(1, f['key']), (2, f['unit']), (4, f['definition']), (5, f['period'])]: text(ws.cell(i, col), value)
            ws.row_dimensions[i].height = 44; ws.cell(i, 4).alignment = Alignment(wrap_text=True, vertical='center')
            if f['role'] == 'input': numeric(ws, keys[f['key']], inputs[f['key']]['value'], f['unit'])
            else:
                raw = expression(parse_formula(f['formula'])[0], refs)
                formula(ws, keys[f['key']], f'=IF(B$3,{raw},NA())', saved['values'][f['key']]['value'])
                ws[keys[f['key']]].number_format = number_format(saved['values'][f['key']]['value'], f['unit'])
        checks = []; start = len(fields) + 10
        header(ws, start, ['Control field', 'Expected', 'Tolerance', 'Reason', 'Passed'])
        for i, c in enumerate(source['checks'], start + 1):
            text(ws.cell(i, 1), c['key']); numeric(ws, f'B{i}', c['expected']); numeric(ws, f'C{i}', c['tolerance'])
            text(ws.cell(i, 4), c['note']); ws.cell(i, 4).alignment = Alignment(wrap_text=True); ws.row_dimensions[i].height = 52
            formula(ws, f'E{i}', f'=IFERROR(AND(ISNUMBER(B{i}),ISNUMBER(C{i}),C{i}>=0,ABS({keys[c["key"]]}-B{i})<=C{i}),FALSE)', True)
            checks.append('E' + str(i))
        formula(ws, 'B4', '=AND(' + ','.join(checks) + ')' if checks else '=TRUE', True)
        text(ws['A5'], 'Controls: ' + ('configured' if checks else 'none; no independent reconciliation'))
    for name, inputs in data['scenarios'].items():
        ws = sheet('Val_' + name, 'Common equity valuation: ' + name)
        fields = data['contract']['fields']; refs = {f['cell']: 'C' + str(i + 8) for i, f in enumerate(fields)}
        keys = {f['key']: refs[f['cell']] for f in fields}; val_refs[name] = keys
        text(ws['A3'], 'Input domain valid'); text(ws['A4'], 'Value positive')
        tests = [f'AND(ISNUMBER({keys[k]}),ABS({keys[k]})<=1E30)' for k in inputs]
        tests += domain(data['method'], keys, data['operating']) + [f"'Calc_{name}'!B4"]
        formula(ws, 'B3', '=IFERROR(AND(' + ','.join(tests) + '),FALSE)', True)
        formula(ws, 'B4', f'=IFERROR({keys["equity_value"]}>0,FALSE)', True)
        text(ws['A5'], 'Blue: manual inputs. Green: source links. See Evidence for exact decimals and provenance.')
        header(ws, 7, ['Field', 'Unit', 'Value / formula', 'Definition', 'Period'])
        for i, f in enumerate(fields, 8):
            for col, value in [(1, f['key']), (2, f['unit']), (4, f['definition']), (5, f['period'])]: text(ws.cell(i, col), value)
            ws.row_dimensions[i].height = 44; ws.cell(i, 4).alignment = Alignment(wrap_text=True, vertical='center')
            cell = keys[f['key']]
            if f['role'] == 'input':
                item = inputs[f['key']]
                if item.get('source_key'):
                    ref = f"'Calc_{name}'!{build_refs[name][item['source_key']]}"
                    formula(ws, cell, f'=IF(ISNUMBER({ref}),{ref},NA())', item['value'])
                else: numeric(ws, cell, item['value'], f['unit'])
            else:
                raw = expression(parse_formula(f['formula'])[0], refs)
                formula(ws, cell, f'=IF(B$3,{raw},NA())', result['scenarios'][name]['values'][f['key']]['value'])
            ws[cell].number_format = number_format(result['scenarios'][name]['values'][f['key']]['value'], f['unit'])
    for i, (name, scenario) in enumerate(result['scenarios'].items(), 8):
        text(summary.cell(i, 1), name.title()); keys = val_refs[name]
        for col, key in [('B', 'per_share'), ('C', 'equity_value')]:
            formula(summary, f'{col}{i}', f'=IF(\'Val_{name}\'!B4,\'Val_{name}\'!{keys[key]},NA())', scenario['values'][key]['value'])
            summary[f'{col}{i}'].number_format = number_format(scenario['values'][key]['value'])
        if data['quote']:
            formula(summary, f'D{i}', f'=IF(AND(ISNUMBER(B14),B14>0),B{i}/B14-1,NA())', scenario['upside']); summary[f'D{i}'].number_format = '0.00%'
        else: text(summary[f'D{i}'], 'Unavailable: no quote')
        formula(summary, f'E{i}', f"='Val_{name}'!B3", True)
        for col in ('B','C','D','E'): summary[f'{col}{i}'].font = Font(name='Arial',color=GREEN,size=11)
    text(summary['A14'], 'Supplied quote (' + data['currency'] + ')')
    if data['quote']: numeric(summary, 'B14', data['quote']['price'])
    else: text(summary['B14'], 'Not supplied')
    notes = [review_label(report), 'Saved receipt: ' + report['id'],
             'Excel uses binary floating point (~15 digits); the server uses Decimal (34 digits).',
             'Numbers are cached from the server, then recalculated by Excel. Exact source decimals are in Evidence.',
             'Editing this file creates an unreviewed local analysis. It does not modify the saved valuation.',
             'Workbook formula errors and FALSE domain/controls require correction; never interpret them as zero.',
             'Sensitivity follows the current assumptions; its steps and multiple factors can be changed on Sensitivity.',
             *data['warnings']]
    for i, note in enumerate(notes, 16):
        summary.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        text(summary.cell(i, 1), note); summary.cell(i, 1).alignment = Alignment(wrap_text=True); summary.row_dimensions[i].height = 32
    sens = sheet('Sensitivity', 'Sensitivity: value per common share (' + data['currency'] + ')')
    sens.column_dimensions['D'].width = 26
    dcf = data['method'] in ('ddm', 'fcfe')
    text(sens['A3'], 'Sensitivity steps' if dcf else 'Multiple factors')
    if dcf:
        numeric(sens, 'B3', result['sensitivity_policy']['discount_step'], 'ratio')
        numeric(sens, 'C3', result['sensitivity_policy']['growth_step'], 'ratio')
        text(sens['B4'], 'Cost of equity step'); text(sens['C4'], 'Terminal growth step')
    else:
        for col, factor in zip(('B','C','D'), result['sensitivity_policy']['multiple_factors']): numeric(sens, col+'3', factor)
    header(sens, 7, ['Scenario', 'Cost of equity / multiple', 'Terminal growth', 'Per share', 'Input domain'])
    row = 8
    for name, scenario in result['scenarios'].items():
        fields = data['contract']['fields']; keys = val_refs[name]
        base = {f['key']: f"'Val_{name}'!{keys[f['key']]}" for f in fields if f['role'] == 'input'}
        for index, point in enumerate(scenario['sensitivity']):
            text(sens.cell(row, 1), name.title()); changes = dict(base)
            for col, (key, value) in zip(('B', 'C'), point['inputs'].items()):
                if dcf:
                    offset = index // 3 - 1 if key == 'cost_equity' else index % 3 - 1
                    axis = f'={base[key]}+{offset}*{col}$3'
                else: axis = f'={base[key]}*{("B","C","D")[index]}$3'
                formula(sens, f'{col}{row}', axis, value)
                sens[f'{col}{row}'].number_format = '0.00%' if dcf else NUM
                changes[key] = f'{col}{row}'
            expressions = {}
            for cell in data['contract']['order']:
                f = next(f for f in fields if f['cell'] == cell)
                expressions[cell] = changes[f['key']] if f['role'] == 'input' else expression(parse_formula(f['formula'])[0], expressions)
            per_share = expressions[next(f['cell'] for f in fields if f['key'] == 'per_share')]
            valid = domain(data['method'], changes, data['operating']) + [f"'Val_{name}'!B3", per_share + '>0']
            valid += [f'ISNUMBER({changes[k]})' for k in point['inputs']]
            valid += ['ISNUMBER(B$3)', 'ISNUMBER(C$3)'] if dcf else [f'ISNUMBER({("B","C","D")[index]}$3)']
            formula(sens, f'E{row}', '=IFERROR(AND(' + ','.join(valid) + '),FALSE)', point['state'] != 'blocked')
            formula(sens, f'D{row}', f'=IF(E{row},{per_share},"Blocked")', point['per_share']['value'] if point['per_share'] else 'Blocked')
            sens[f'D{row}'].number_format = number_format(point['per_share']['value']) if point['per_share'] else NUM; row += 1
    proof = sheet('Evidence', 'Exact inputs, assumptions and provenance')
    header(proof, 7, ['Section', 'Item', 'Exact text (continuations retain all content)', '', ''])
    proof.column_dimensions['A'].width = 32; proof.column_dimensions['B'].width = 16
    proof.column_dimensions['C'].width = 100; proof.column_dimensions['D'].width = 2; proof.column_dimensions['E'].width = 2
    row = 8
    for title, lines in sections(report):
        for index, line in enumerate(lines, 1):
            # Short chunks keep every source character visible rather than clipping a 32k cell.
            line = clean(line)
            for offset in range(0, max(1, len(line)), 220):
                text(proof.cell(row, 1), title); text(proof.cell(row, 2), str(index) + (' cont.' if offset else ''))
                text(proof.cell(row, 3), line[offset:offset + 220])
                proof.cell(row, 3).alignment = Alignment(wrap_text=True, vertical='top'); proof.row_dimensions[row].height = 54; row += 1
    for ws in wb:
        ws.print_options.horizontalCentered = True
        ws.print_area = f'A1:E{ws.max_row}'
    stream = BytesIO(); wb.save(stream)
    return cached_values(stream.getvalue(), wb.sheetnames, caches)


def cached_values(raw, sheetnames, caches):
    """OOXML caches reproduce the saved receipt; auto-calc remains enabled for edits."""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    out = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(out, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            blob = source.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                index = int(item.filename.split('sheet')[-1].split('.')[0]) - 1
                tree = ET.fromstring(blob)
                for cell in tree.iter(ns + 'c'):
                    key = (sheetnames[index], cell.attrib['r'])
                    if key not in caches: continue
                    value = caches[key]; node = cell.find(ns + 'v')
                    if node is None: node = ET.SubElement(cell, ns + 'v')
                    if isinstance(value, bool): cell.set('t', 'b'); node.text = '1' if value else '0'
                    elif value == 'Blocked': cell.set('t', 'str'); node.text = value
                    else: cell.set('t', 'n'); node.text = str(value)
                blob = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
            target.writestr(item, blob)
    return out.getvalue()
