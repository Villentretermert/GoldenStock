"""Exercise compatibility, immutable versions and fund boundaries via the HTTP API."""
import io
import zipfile
from pathlib import Path
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties
from test_workspace import workspace
from goldenstock.template_contract import validate_template,clean_workbook,formula_dependencies

DEMO=Path('src/goldenstock/assets/demo-template.xlsx').read_bytes()


def modify(change):
    book=load_workbook(io.BytesIO(DEMO));change(book)
    result=io.BytesIO();book.save(result);book.close();return result.getvalue()


def test_demo_and_clean_structure():
    result=validate_template(DEMO)
    assert len(result['fields'])==7
    clean=clean_workbook(result)
    assert validate_template(clean)==result
    wb=load_workbook(io.BytesIO(clean))
    assert wb['Model']['B11'].value is None
    assert wb['Model']['F7'].value=='=E7*(1+$B$12)'


@pytest.mark.parametrize('formula',['=B7+1','=C7+1','=SUM(B11:B12)',"='[private.xlsx]Sheet1'!B1",'=INDIRECT("B11")','=B11+999','=B11+v0','=B300'])
def test_cycles_functions_links_constants_and_unmapped_refs_rejected(formula):
    content=modify(lambda b:setattr(b['Model']['B7'],'value',formula))
    with pytest.raises(HTTPException) as err: validate_template(content)
    assert err.value.status_code==422


@pytest.mark.parametrize('change',[
    lambda b:setattr(b['Model']['B11'],'value','=1+1'),
    lambda b:setattr(b['Model']['B12'],'value','private string'),
    lambda b:setattr(b['Model']['B11'],'hyperlink','https://example.com'),
    lambda b:setattr(b['Model'],'sheet_state','hidden'),
    lambda b:setattr(b,'calculation',CalcProperties(iterate=True)),
    lambda b:setattr(b['GS_Map']['A3'],'value','profit_owners'),
    lambda b:setattr(b['GS_Map']['F2'],'value','=WEBSERVICE("https://example.com")'),
])
def test_unsupported_workbook_features(change):
    with pytest.raises(HTTPException): validate_template(modify(change))


def test_vba_part_rejected():
    stream=io.BytesIO(DEMO)
    with zipfile.ZipFile(stream,'a') as z:z.writestr('xl/vbaProject.bin',b'synthetic')
    with pytest.raises(HTTPException): validate_template(stream.getvalue())


def post(client,content=DEMO,template_id='',key='template-test-01'):
    return client.post('/api/v1/templates',headers={'Idempotency-Key':key},
        data={'title':'Synthetic five-year template','template_id':template_id},
        files={'file':('demo.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})


def test_template_privacy_versions_copies_and_redaction(workspace):
    _,connect=workspace;a,b,e=connect(),connect('bob'),connect('eve')
    def secrets(book):
        book['Model']['B11']=987654321
        book['Model']['A20']='Private analyst note'
        book['Model']['B12'].comment=Comment('Private comment','Private author')
        book.properties.creator='Private author'
    original=modify(secrets)
    response=post(a,original);assert response.status_code==201,response.text
    version=response.json();tid=version['template_id'];vid=version['id']
    assert post(a,original).json()['id']==vid
    assert post(a,DEMO).status_code==409
    assert b.get('/api/v1/templates/'+tid).status_code==404
    assert e.get('/api/v1/templates/'+tid).status_code==404
    assert a.patch('/api/v1/templates/'+tid+'/visibility',json={'shared':True,'revision':1}).status_code==200
    assert b.get('/api/v1/templates/'+tid).status_code==200
    assert e.get('/api/v1/templates/'+tid).status_code==404
    url=f'/api/v1/templates/{tid}/versions/{vid}/'
    assert a.get(url+'original').content==original
    assert b.get(url+'original').status_code==404
    assert b.patch('/api/v1/templates/'+tid+'/visibility',json={'shared':False,'revision':2}).status_code==404
    clean=b.get(url+'structure').content
    wb=load_workbook(io.BytesIO(clean))
    assert wb['Model']['B11'].value is None and wb['Model']['A20'].value is None
    assert wb['Model']['B12'].comment is None and wb.properties.creator!='Private author'
    copied=b.post('/api/v1/templates/'+tid+'/copy').json()
    assert copied['owner_id']!=a.get('/api/v1/templates/'+tid).json()['owner_id']
    assert a.get('/api/v1/templates/'+copied['id']).status_code==404
    assert post(b,template_id=tid,key='template-write-denied').status_code==404
    v2=post(a,template_id=tid,key='template-second-01').json()
    assert v2['number']==2
    assert b.get(url+'structure').status_code==404
    assert a.get(url+'original').content==original
    assert b.get('/api/v1/templates/'+copied['id']).status_code==200
    assert a.patch('/api/v1/templates/'+tid+'/visibility',json={'shared':True,'revision':2}).status_code==409


def test_template_csrf_and_ui(workspace):
    app,connect=workspace;a=connect()
    from goldenstock.ui import mount_ui
    mount_ui(app)
    assert a.get('/api/v1/templates/demo').content==DEMO
    page=a.get('/templates');assert 'Validate and save template' in page.text
    row=post(a).json()
    assert 'Not calculated' in a.get('/templates/'+row['template_id']).text
    a.headers['x-csrf-token']='wrong'
    assert post(a,key='template-csrf-test').status_code==403
