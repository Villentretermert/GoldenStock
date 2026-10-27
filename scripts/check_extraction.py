"""Run via stdin inside the local worker container; no customer data is used."""
import tempfile
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from goldenstock.parser_process import run_parser
from goldenstock.extraction import ExtractionError

with tempfile.TemporaryDirectory() as directory:
    root=Path(directory)
    image=Image.new('RGB',(1200,400),'white')
    font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',40)
    ImageDraw.Draw(image).text((60,80),'Profit -120 USD million',font=font,fill='black')
    scan=root/'scan.pdf'; image.save(scan,'PDF',resolution=150)
    result=run_parser(scan,'.pdf',lambda:True)
    words=' '.join(item['text'] for item in result['items'])
    assert 'Profit' in words and '-120' in words and 'million' in words, words
    assert all(item['method']=='ocr' and item['location']['page']==1 for item in result['items'])
    assert all(0<=item['confidence']<=1 for item in result['items'])
    book=Workbook(); book.active.append(['Profit',-120]); book.active.merge_cells('A2:B2')
    xlsx=root/'sample.xlsx'; book.save(xlsx)
    result=run_parser(xlsx,'.xlsx',lambda:True)
    assert result['items'][1]['text']=='-120'
    assert result['items'][3]['location']['merged_range']=='A2:B2'
    deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6])
    table=slide.shapes.add_table(1,2,Inches(1),Inches(1),Inches(4),Inches(1)).table
    table.cell(0,0).text='Profit'; table.cell(0,1).text='(120)'
    pptx=root/'sample.pptx'; deck.save(pptx)
    result=run_parser(pptx,'.pptx',lambda:True)
    assert result['items'][1]['text']=='(120)'
    try:
        run_parser(scan,'.pdf',lambda:True,timeout=.001)
        raise AssertionError('Timeout was not enforced')
    except ExtractionError as exc:
        assert str(exc)=='extraction_timeout'
print('PASS: real local OCR, XLSX merged cells, PPTX cells, bounded Linux subprocess and timeout. Synthetic corpus only.')
